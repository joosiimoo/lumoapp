"""Daily sales export: one confirmed-sale projection, two serializers, no writes."""

from __future__ import annotations

import csv
import io
import inspect
from datetime import UTC, datetime, time, timedelta
from decimal import Decimal
from unittest.mock import MagicMock
from uuid import UUID
from zoneinfo import ZoneInfo

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select, text
from sqlalchemy.dialects import postgresql

from app.agent.generative_ui import GenerativeUIRegistry
from app.agent.providers.scripted import ScriptedLLMProvider
from app.agent.registrations import (
    register_conversational_sale_tools,
    register_daily_close_confirmed_ui,
    register_daily_close_preparation_ui,
    register_operational_day_summary_ui,
    register_sale_confirmed_ui,
    register_sale_item_added_ui,
    register_sale_summary_ui,
)
from app.agent.tools import ToolRegistry
from app.agent.ui_actions import UiActionRegistry
from app.application.queries.export_daily_sales import (
    SalesExport,
    SalesExportLine,
    business_slug,
)
from app.domain.shared.ids import new_uuid7
from app.domain.shared.tenant import TenantContext
from app.infrastructure.export.columns import COLUMN_WIDTHS, EXPORT_COLUMNS
from app.infrastructure.export.daily_sales_csv import render_daily_sales_csv
from app.infrastructure.export.daily_sales_xlsx import render_daily_sales_xlsx
from app.infrastructure.persistence.models import (
    AuditEventRow,
    CashCountRow,
    ClosingSnapshotRow,
    IdempotencyRecordRow,
    OperationalDayRow,
    OutboxEventRow,
    PaymentRow,
    ProductRow,
    SaleItemRow,
    SaleSessionRow,
)
from app.infrastructure.persistence.rls import set_current_business_id
from app.infrastructure.persistence.sales_export import SalesExportRepository
from tests.conftest import seed_business
from tests.sale_cleanup import clear_tenant_sale_mutations

ZONE = ZoneInfo("America/Mexico_City")
SID_MIXED = UUID("11111111-1111-4111-8111-111111111111")
SID_CARD = UUID("00000000-0000-4000-8000-000000000001")
SID_TRANSFER = UUID("00000000-0000-4000-8000-000000000002")
ITEM_PINEAPPLE = UUID("22222222-2222-4222-8222-222222222222")
ITEM_OVERRIDE = UUID("00000000-0000-4000-8000-0000000000aa")
ITEM_FREE = UUID("33333333-3333-4333-8333-333333333333")
UNSUPPORTED = (
    "Puedo registrar un producto del catálogo con cantidad y unidad. Prueba con *900gr zanahoria*."
)


def test_slug_filename_and_serializers_are_deterministic() -> None:
    assert business_slug("Carrota") == "carrota"
    assert business_slug("Ñandú & Hijos") == "nandu-hijos"
    assert business_slug("!!!") == "negocio"
    assert business_slug("a-" * 30) == ("a-" * 20).strip("-")
    assert '"' not in business_slug('Say "Hi"')
    assert "/" not in business_slug("Café / Norte")
    assert ".." not in business_slug("..")

    export = _sample_export()
    first_csv = render_daily_sales_csv(export)
    second_csv = render_daily_sales_csv(export)
    assert first_csv == second_csv
    assert first_csv.startswith(b"\xef\xbb\xbf")
    assert first_csv.endswith(b"\r\n")
    assert b"\r\n" in first_csv
    text_body = first_csv.decode("utf-8-sig")
    assert '"Piña, fresca"' in text_body
    assert 'Pan ""bolillo""' in text_body
    rows = list(csv.reader(io.StringIO(text_body)))
    assert rows[0] == list(EXPORT_COLUMNS)
    assert rows[1][4] == "Piña, fresca"
    assert rows[1][7] == "0.9"
    assert rows[1][9] == "20.00"
    assert rows[1][11] == ""
    assert rows[2][11] == "precio de feria"
    assert rows[3][5] == "free_concept"
    assert rows[3][6] == ""
    assert rows[3][9] == ""
    assert rows[3][4] == "Nota\nextra"
    assert rows[1][2] == "2026-09-23T14:05:00-06:00"
    assert rows[1][16] == "49.00"
    assert rows[2][16] == "49.00"
    assert rows[3][16] == "49.00"

    first_xlsx = render_daily_sales_xlsx(export)
    second_xlsx = render_daily_sales_xlsx(export)
    assert _workbook_facts(first_xlsx) == _workbook_facts(second_xlsx)
    facts = _workbook_facts(first_xlsx)
    assert facts["sheets"] == ["Ventas"]
    assert facts["header"] == list(EXPORT_COLUMNS)
    assert facts["freeze"] == "A2"
    assert facts["widths"] == list(COLUMN_WIDTHS)
    assert facts["money_formats"] == ["0.00"]
    assert facts["quantity_format"] == "0.######"
    assert facts["date_format"] == "yyyy-mm-dd"
    assert facts["values"][0][4] == "Piña, fresca"
    assert facts["values"][0][7] == pytest.approx(0.9)
    assert facts["values"][0][9] == pytest.approx(20.0)
    assert facts["values"][2][6] is None
    assert facts["values"][2][9] is None
    assert facts["values"][0][2] == datetime(2026, 9, 23, 14, 5, 0)
    assert facts["formulas"] == []

    source = inspect.getsource(__import__("app.application.queries.export_daily_sales", fromlist=["x"]))
    assert "openpyxl" not in source
    assert "fastapi" not in source.lower()
    assert "FastAPI" not in source


def test_export_sql_is_one_correlated_select() -> None:
    session = MagicMock()
    SalesExportRepository(session).load(
        tenant=TenantContext(business_id=UUID("aaaaaaaa-aaaa-4aaa-8aaa-aaaaaaaaaaaa"), actor_id=new_uuid7()),
        day_id=UUID("bbbbbbbb-bbbb-4bbb-8bbb-bbbbbbbbbbbb"),
    )
    statements = [call.args[0] for call in session.execute.call_args_list]
    assert len(statements) == 2
    statement = statements[-1]
    sql = str(statement.compile(dialect=postgresql.dialect()))
    assert sql.upper().count("SELECT") >= 1
    assert "FOR UPDATE" not in sql.upper()
    assert "LOCK TABLE" not in sql.upper()
    assert "current_price" not in sql
    assert "closing_snapshot" not in sql.lower()
    assert "s.operational_day_id = operations.operational_days.id" in sql
    assert "sales.sale_items.product_name_snapshot" in sql


def test_export_is_not_a_tool_or_ui_action() -> None:
    tools = ToolRegistry()
    register_conversational_sale_tools(tools)
    assert tools.is_registered("operational_day.export_sales@1") is False
    assert tools.is_registered("operational_day.summary@1") is True
    ui = GenerativeUIRegistry()
    register_sale_item_added_ui(ui)
    register_sale_summary_ui(ui)
    register_sale_confirmed_ui(ui)
    register_operational_day_summary_ui(ui)
    register_daily_close_preparation_ui(ui)
    register_daily_close_confirmed_ui(ui)
    assert ui.get("export_ready_card", 1) is None
    actions = UiActionRegistry()
    assert actions.ids() == [
        "sale.pay.cash@1",
        "sale.pay.card@1",
        "sale.pay.transfer@1",
        "closing.request@1",
        "closing.confirm@1",
    ]
    assert actions.is_registered("operational_day.export_sales@1") is False
    provider = ScriptedLLMProvider()
    for phrase in (
        "exporta las ventas de hoy",
        "exportar ventas",
        "descargar excel",
        "descargar csv",
    ):
        decision = provider.interpret(phrase, {}, [])
        assert decision.intent == "unsupported"
        assert decision.candidate_tool is None
        assert decision.clarification_question == UNSUPPORTED


def test_open_day_csv_and_xlsx_match_persisted_sales(client: TestClient, db_session) -> None:
    business_id, user_id, token = seed_business(db_session, name="Ñandú & Hijos")
    try:
        business_date = datetime.now(UTC).astimezone(ZONE).date()
        local_confirmed = datetime.combine(business_date, time(14, 5), tzinfo=ZONE)
        confirmed_at = local_confirmed.astimezone(UTC)
        day_id = _day(db_session, business_id, business_date=business_date)
        pineapple_id = _product(db_session, business_id, name="Piña", price="20.00")
        bread_id = _product(db_session, business_id, name="Bolillo", price="20.00")
        tomato_id = _product(db_session, business_id, name="Jitomate", price="10.00")
        milk_id = _product(db_session, business_id, name="Leche", price="25.00")
        _sale(
            db_session,
            business_id=business_id,
            actor_id=user_id,
            day_id=day_id,
            session_id=SID_MIXED,
            confirmed_at=confirmed_at,
            method="cash",
            amount="49.00",
            items=[
                _item(
                    item_id=ITEM_PINEAPPLE,
                    product_id=pineapple_id,
                    name="Piña, fresca",
                    quantity="0.900",
                    unit="kilogram",
                    catalog="20.00",
                    price="20.00",
                    total="18.00",
                    created_at=confirmed_at + timedelta(seconds=1),
                ),
                _item(
                    item_id=ITEM_OVERRIDE,
                    product_id=bread_id,
                    name='Pan "bolillo"',
                    quantity="1",
                    unit="unit",
                    catalog="20.00",
                    price="15.00",
                    total="15.00",
                    reason="precio de feria",
                    created_at=confirmed_at + timedelta(seconds=2),
                ),
                _item(
                    item_id=ITEM_FREE,
                    product_id=None,
                    name="Nota\nextra",
                    source_type="free_concept",
                    quantity="2",
                    unit="unit",
                    catalog=None,
                    price="8.00",
                    total="16.00",
                    created_at=confirmed_at + timedelta(seconds=3),
                ),
            ],
        )
        card_item = new_uuid7()
        transfer_item = new_uuid7()
        _sale(
            db_session,
            business_id=business_id,
            actor_id=user_id,
            day_id=day_id,
            session_id=SID_CARD,
            confirmed_at=confirmed_at + timedelta(hours=1),
            method="card",
            amount="10.00",
            items=[
                _item(
                    item_id=card_item,
                    product_id=tomato_id,
                    name="Jitomate",
                    quantity="1",
                    unit="unit",
                    catalog="10.00",
                    price="10.00",
                    total="10.00",
                    created_at=confirmed_at + timedelta(hours=1),
                )
            ],
        )
        _sale(
            db_session,
            business_id=business_id,
            actor_id=user_id,
            day_id=day_id,
            session_id=SID_TRANSFER,
            confirmed_at=confirmed_at + timedelta(hours=2),
            method="transfer",
            amount="25.00",
            items=[
                _item(
                    item_id=transfer_item,
                    product_id=milk_id,
                    name="Leche",
                    quantity="1",
                    unit="package",
                    catalog="25.00",
                    price="25.00",
                    total="25.00",
                    created_at=confirmed_at + timedelta(hours=2),
                )
            ],
        )
        _open_session(db_session, business_id, user_id, pineapple_id, "NO-OPEN", "open", "conv-open")
        _open_session(
            db_session, business_id, user_id, pineapple_id, "NO-READY", "ready_to_charge", "conv-ready"
        )
        pineapple = db_session.get(ProductRow, pineapple_id)
        pineapple.current_price = Decimal("99.00")
        pineapple.name = "Renombrada"
        db_session.commit()

        before = _write_counts(db_session, business_id)
        csv_response = _export(client, token, str(day_id), "csv")
        csv_again = _export(client, token, "current", "csv")
        xlsx_response = _export(client, token, str(day_id), "xlsx")
        xlsx_again = _export(client, token, str(day_id), "xlsx")
        db_session.expire_all()
        assert _write_counts(db_session, business_id) == before
        assert db_session.get(ProductRow, pineapple_id).current_price == Decimal("99.00")
        assert db_session.get(OperationalDayRow, day_id).status == "open"

        assert csv_response.status_code == 200, csv_response.text
        assert csv_response.headers["cache-control"] == "no-store"
        assert "text/csv" in csv_response.headers["content-type"]
        assert "charset=utf-8" in csv_response.headers["content-type"]
        filename = f"lumo-nandu-hijos-ventas-{business_date.isoformat()}.csv"
        assert csv_response.headers["content-disposition"] == f'attachment; filename="{filename}"'
        assert str(day_id) not in filename
        assert csv_response.content == csv_again.content
        rows = _csv_rows(csv_response.content)
        assert [row[3] for row in rows] == [
            str(ITEM_PINEAPPLE),
            str(ITEM_OVERRIDE),
            str(ITEM_FREE),
            str(card_item),
            str(transfer_item),
        ]
        assert rows[0][2] == local_confirmed.isoformat(timespec="seconds")
        assert rows[0][4] == "Piña, fresca"
        assert rows[0][9] == "20.00"
        assert "NO-OPEN" not in csv_response.text
        assert "NO-READY" not in csv_response.text
        assert "Renombrada" not in csv_response.text
        assert {row[15] for row in rows} == {"cash", "card", "transfer"}
        _assert_reconciliation(rows)

        assert xlsx_response.status_code == 200
        assert xlsx_response.headers["content-type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert xlsx_response.headers["content-disposition"].endswith('.xlsx"')
        first_facts = _workbook_facts(xlsx_response.content)
        second_facts = _workbook_facts(xlsx_again.content)
        assert first_facts == second_facts
        assert first_facts["sheets"] == ["Ventas"]
        assert len(first_facts["values"]) == 5
        assert first_facts["values"][0][9] == pytest.approx(20.0)
        assert isinstance(first_facts["values"][0][7], (int, float))
        assert isinstance(first_facts["values"][0][12], (int, float))

        _close_with_empty_snapshot(db_session, business_id, user_id, day_id, business_date)
        closed = _export(client, token, str(day_id), "csv")
        assert closed.status_code == 200, closed.text
        assert _csv_rows(closed.content) == rows
        set_current_business_id(db_session, business_id)
        db_session.expire_all()
        snapshot = db_session.scalars(
            select(ClosingSnapshotRow).where(ClosingSnapshotRow.operational_day_id == day_id)
        ).one()
        assert snapshot.sale_count == 0
        assert db_session.get(OperationalDayRow, day_id).status == "closed"
    finally:
        db_session.rollback()
        clear_tenant_sale_mutations(db_session, business_id)


def test_empty_existing_day_is_header_only(client: TestClient, db_session) -> None:
    business_id, _, token = seed_business(db_session, name="Vacío")
    try:
        business_date = datetime.now(UTC).astimezone(ZONE).date()
        day_id = _day(db_session, business_id, business_date=business_date)
        db_session.commit()
        response = _export(client, token, "current", "csv")
        workbook = _export(client, token, str(day_id), "xlsx")
        assert response.status_code == 200, response.text
        assert _csv_rows(response.content) == []
        assert response.content.decode("utf-8-sig").split("\r\n", 1)[0] == ",".join(EXPORT_COLUMNS)
        facts = _workbook_facts(workbook.content)
        assert facts["values"] == []
        assert facts["header"] == list(EXPORT_COLUMNS)
    finally:
        db_session.rollback()
        clear_tenant_sale_mutations(db_session, business_id)


def test_missing_foreign_and_invalid_refs(client: TestClient, db_session) -> None:
    business_id, _, token = seed_business(db_session, name="Dueño")
    other_id, _, other_token = seed_business(db_session, name="Ajeno")
    try:
        day_id = _day(db_session, business_id, business_date=datetime.now(UTC).astimezone(ZONE).date())
        _product_sale(db_session, business_id, day_id)
        db_session.commit()
        missing = _export(client, other_token, "current", "csv")
        foreign = _export(client, other_token, str(day_id), "xlsx")
        unknown = _export(client, token, str(new_uuid7()), "csv")
        for response in (missing, foreign, unknown):
            assert response.status_code == 404
            assert response.json()["error"]["code"] == "TENANT_SCOPE_VIOLATION"
            assert response.json()["error"]["message"] == "operational day not found"
            assert not response.content.startswith(b"\xef\xbb\xbf")
            assert not response.content.startswith(b"PK")
        invalid_ref = _export(client, token, "today", "csv")
        invalid_format = client.get(
            f"/api/v1/operational-days/{day_id}/sales-export?format=CSV",
            headers={"Authorization": f"Bearer {token}"},
        )
        missing_format = client.get(
            f"/api/v1/operational-days/{day_id}/sales-export",
            headers={"Authorization": f"Bearer {token}"},
        )
        for response in (invalid_ref, invalid_format, missing_format):
            assert response.status_code == 422
            assert response.json()["error"]["code"] == "VALIDATION_ERROR"
        set_current_business_id(db_session, other_id)
        hidden = db_session.execute(
            text("SELECT id FROM sales.sale_items WHERE business_id = :business_id"),
            {"business_id": business_id},
        ).all()
        assert hidden == []
    finally:
        db_session.rollback()
        clear_tenant_sale_mutations(db_session, business_id)
        clear_tenant_sale_mutations(db_session, other_id)


def test_inconsistent_sale_and_invalid_timezone_return_no_file(client: TestClient, db_session) -> None:
    business_id, user_id, token = seed_business(db_session, name="Roto")
    try:
        day_id = _day(db_session, business_id, business_date=datetime(2026, 1, 15).date())
        product_id = _product(db_session, business_id, name="Pan", price="10.00")
        _sale(
            db_session,
            business_id=business_id,
            actor_id=user_id,
            day_id=day_id,
            session_id=new_uuid7(),
            confirmed_at=datetime(2026, 1, 15, 18, tzinfo=UTC),
            method="cash",
            amount="10.00",
            items=[
                _item(
                    item_id=new_uuid7(),
                    product_id=product_id,
                    name="Visible",
                    quantity="1",
                    unit="unit",
                    catalog="10.00",
                    price="10.00",
                    total="10.00",
                    created_at=datetime(2026, 1, 15, 18, tzinfo=UTC),
                )
            ],
        )
        broken_session = new_uuid7()
        db_session.add(
            SaleSessionRow(
                id=broken_session,
                business_id=business_id,
                actor_id=user_id,
                status="confirmed",
                currency="MXN",
                operational_day_id=day_id,
                confirmed_at=datetime(2026, 1, 15, 19, tzinfo=UTC),
            )
        )
        db_session.flush()
        db_session.add(
            SaleItemRow(
                business_id=business_id,
                sale_session_id=broken_session,
                product_id=product_id,
                source_type="catalog",
                product_name_snapshot="Sin pago",
                quantity_input=Decimal("1"),
                unit_input="unit",
                quantity_normalized=Decimal("1"),
                unit_normalized="unit",
                unit_price=Decimal("10.00"),
                currency="MXN",
                line_total=Decimal("10.00"),
                catalog_unit_price_snapshot=Decimal("10.00"),
            )
        )
        bad_day = _day(
            db_session,
            business_id,
            business_date=datetime(2026, 1, 16).date(),
            timezone_name="Not/AZone",
        )
        mismatch_day = _day(db_session, business_id, business_date=datetime(2026, 1, 17).date())
        mismatch_session = new_uuid7()
        db_session.add(
            SaleSessionRow(
                id=mismatch_session,
                business_id=business_id,
                actor_id=user_id,
                status="confirmed",
                currency="MXN",
                operational_day_id=mismatch_day,
                confirmed_at=datetime(2026, 1, 17, 18, tzinfo=UTC),
            )
        )
        db_session.flush()
        db_session.add(
            SaleItemRow(
                business_id=business_id,
                sale_session_id=mismatch_session,
                product_id=product_id,
                source_type="catalog",
                product_name_snapshot="Dólar",
                quantity_input=Decimal("1"),
                unit_input="unit",
                quantity_normalized=Decimal("1"),
                unit_normalized="unit",
                unit_price=Decimal("10.00"),
                currency="USD",
                line_total=Decimal("10.00"),
                catalog_unit_price_snapshot=Decimal("10.00"),
            )
        )
        db_session.add(
            PaymentRow(
                business_id=business_id,
                sale_session_id=mismatch_session,
                actor_id=user_id,
                method="cash",
                amount=Decimal("10.00"),
                currency="USD",
                status="recorded",
                source="manual_capture",
            )
        )
        db_session.commit()

        broken = _export(client, token, str(day_id), "csv")
        assert broken.status_code == 500
        assert broken.json()["error"]["code"] == "INTERNAL_ERROR"
        assert broken.json()["error"]["message"] == "confirmed sales export is inconsistent"
        assert b"Visible" not in broken.content
        assert not broken.content.startswith(b"\xef\xbb\xbf")

        invalid_zone = _export(client, token, str(bad_day), "xlsx")
        assert invalid_zone.status_code == 500
        assert invalid_zone.json()["error"]["code"] == "INTERNAL_ERROR"
        assert not invalid_zone.content.startswith(b"PK")

        mismatch = _export(client, token, str(mismatch_day), "csv")
        assert mismatch.status_code == 422
        assert mismatch.json()["error"]["code"] == "VALIDATION_ERROR"
        assert not mismatch.content.startswith(b"\xef\xbb\xbf")
    finally:
        db_session.rollback()
        clear_tenant_sale_mutations(db_session, business_id)


def test_export_phrases_do_not_sell_or_summarize(client: TestClient, db_session) -> None:
    business_id, _, token = seed_business(db_session, name="Frases")
    try:
        for index, phrase in enumerate(
            ("exporta las ventas de hoy", "exportar ventas", "descargar excel", "descargar csv")
        ):
            response = client.post(
                "/api/v1/lumo/messages",
                json={"message": phrase, "conversation_id": "conv-export"},
                headers={
                    "Authorization": f"Bearer {token}",
                    "Idempotency-Key": f"export-phrase-{index}",
                },
            )
            assert response.status_code == 200, response.text
            body = response.json()
            assert body["ui"] == []
            assert body["text"] == UNSUPPORTED
        set_current_business_id(db_session, business_id)
        db_session.expire_all()
        assert db_session.scalars(select(SaleSessionRow)).all() == []
        assert db_session.scalars(select(PaymentRow)).all() == []
        assert db_session.scalars(select(OperationalDayRow)).all() == []
    finally:
        db_session.rollback()
        clear_tenant_sale_mutations(db_session, business_id)


def _sample_export() -> SalesExport:
    confirmed = datetime(2026, 9, 23, 20, 5, tzinfo=UTC)
    payment = UUID("44444444-4444-4444-8444-444444444444")
    session = UUID("55555555-5555-4555-8555-555555555555")

    def line(**kwargs) -> SalesExportLine:
        base = dict(
            business_date=datetime(2026, 9, 23).date(),
            sale_session_id=session,
            sale_confirmed_at=confirmed,
            source_type="catalog",
            quantity=Decimal("1"),
            unit="unit",
            catalog_unit_price=Decimal("20.00"),
            unit_price=Decimal("20.00"),
            price_override_reason=None,
            line_total=Decimal("20.00"),
            currency="MXN",
            payment_id=payment,
            payment_method="cash",
            payment_amount=Decimal("49.00"),
        )
        base.update(kwargs)
        return SalesExportLine(**base)

    return SalesExport(
        business_name="Ñandú & Hijos",
        business_date=datetime(2026, 9, 23).date(),
        timezone_name="America/Mexico_City",
        currency="MXN",
        status="open",
        lines=(
            line(
                sale_item_id=ITEM_PINEAPPLE,
                product_name="Piña, fresca",
                product_id=UUID("66666666-6666-4666-8666-666666666666"),
                quantity=Decimal("0.900"),
                unit="kilogram",
                line_total=Decimal("18.00"),
            ),
            line(
                sale_item_id=ITEM_OVERRIDE,
                product_name='Pan "bolillo"',
                product_id=UUID("77777777-7777-4777-8777-777777777777"),
                unit_price=Decimal("15.00"),
                price_override_reason="precio de feria",
                line_total=Decimal("15.00"),
            ),
            line(
                sale_item_id=ITEM_FREE,
                product_name="Nota\nextra",
                source_type="free_concept",
                product_id=None,
                quantity=Decimal("2"),
                catalog_unit_price=None,
                unit_price=Decimal("8.00"),
                line_total=Decimal("16.00"),
            ),
        ),
    )


def _workbook_facts(payload: bytes) -> dict:
    workbook = load_workbook(io.BytesIO(payload))
    sheet = workbook["Ventas"]
    header = [cell.value for cell in sheet[1]]
    values = []
    formulas = []
    money_formats = set()
    quantity_format = None
    date_format = None
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=17):
        if all(cell.value is None for cell in row):
            continue
        values.append([cell.value for cell in row])
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formulas.append(cell.value)
        if row[7].value is not None:
            quantity_format = row[7].number_format
        if row[0].value is not None:
            date_format = row[0].number_format
        for index in (9, 10, 12, 16):
            if row[index].value is not None:
                money_formats.add(row[index].number_format)
    return {
        "sheets": workbook.sheetnames,
        "header": header,
        "freeze": sheet.freeze_panes,
        "filter": sheet.auto_filter.ref,
        "widths": [sheet.column_dimensions[chr(64 + index)].width for index in range(1, 18)],
        "values": values,
        "formulas": formulas,
        "money_formats": sorted(money_formats),
        "quantity_format": quantity_format,
        "date_format": date_format,
    }


def _csv_rows(payload: bytes) -> list[list[str]]:
    rows = list(csv.reader(io.StringIO(payload.decode("utf-8-sig"))))
    assert rows[0] == list(EXPORT_COLUMNS)
    return rows[1:]


def _assert_reconciliation(rows: list[list[str]]) -> None:
    line_total = sum(Decimal(row[12]) for row in rows)
    naive_payment = sum(Decimal(row[16]) for row in rows)
    distinct: dict[str, Decimal] = {}
    for row in rows:
        distinct[row[14]] = Decimal(row[16])
    assert line_total == sum(distinct.values())
    assert naive_payment != line_total
    grouped: dict[str, set[str]] = {}
    for row in rows:
        grouped.setdefault(row[1], set()).add(row[16])
    assert all(len(amounts) == 1 for amounts in grouped.values())


def _export(client: TestClient, token: str, ref: str, export_format: str):
    return client.get(
        f"/api/v1/operational-days/{ref}/sales-export?format={export_format}",
        headers={"Authorization": f"Bearer {token}"},
    )


def _write_counts(session, business_id) -> tuple[int, int, int]:
    set_current_business_id(session, business_id)
    session.expire_all()

    def count(model) -> int:
        return session.scalar(select(func.count()).select_from(model).where(model.business_id == business_id))

    return count(AuditEventRow), count(OutboxEventRow), count(IdempotencyRecordRow)


def _day(session, business_id, *, business_date, timezone_name: str = "America/Mexico_City") -> UUID:
    set_current_business_id(session, business_id)
    day_id = new_uuid7()
    session.add(
        OperationalDayRow(
            id=day_id,
            business_id=business_id,
            business_date=business_date,
            status="open",
            timezone=timezone_name,
        )
    )
    session.flush()
    return day_id


def _product(session, business_id, *, name: str, price: str) -> UUID:
    set_current_business_id(session, business_id)
    product_id = new_uuid7()
    session.add(
        ProductRow(
            id=product_id,
            business_id=business_id,
            name=name,
            normalized_name=name.lower(),
            sale_unit="unit",
            pricing_type="per_unit",
            current_price=Decimal(price),
            status="active",
        )
    )
    session.flush()
    return product_id


def _item(**kwargs) -> dict:
    return kwargs


def _sale(session, *, business_id, actor_id, day_id, session_id, confirmed_at, method, amount, items) -> None:
    set_current_business_id(session, business_id)
    session.add(
        SaleSessionRow(
            id=session_id,
            business_id=business_id,
            actor_id=actor_id,
            status="confirmed",
            currency="MXN",
            operational_day_id=day_id,
            confirmed_at=confirmed_at,
        )
    )
    session.flush()
    for item in items:
        session.add(
            SaleItemRow(
                id=item["item_id"],
                business_id=business_id,
                sale_session_id=session_id,
                product_id=item["product_id"],
                source_type=item.get("source_type", "catalog"),
                product_name_snapshot=item["name"],
                quantity_input=Decimal(item["quantity"]),
                unit_input=item["unit"] if item["unit"] != "kilogram" else "gram",
                quantity_normalized=Decimal(item["quantity"]),
                unit_normalized=item["unit"],
                unit_price=Decimal(item["price"]),
                currency="MXN",
                line_total=Decimal(item["total"]),
                catalog_unit_price_snapshot=None if item["catalog"] is None else Decimal(item["catalog"]),
                price_override_reason=item.get("reason"),
                created_at=item["created_at"],
            )
        )
    session.add(
        PaymentRow(
            business_id=business_id,
            sale_session_id=session_id,
            actor_id=actor_id,
            method=method,
            amount=Decimal(amount),
            currency="MXN",
            status="recorded",
            source="manual_capture",
        )
    )
    session.flush()


def _open_session(session, business_id, actor_id, product_id, name: str, status: str, conversation_id: str) -> None:
    set_current_business_id(session, business_id)
    session_id = new_uuid7()
    session.add(
        SaleSessionRow(
            id=session_id,
            business_id=business_id,
            actor_id=actor_id,
            conversation_id=conversation_id,
            status=status,
            currency="MXN",
        )
    )
    session.flush()
    session.add(
        SaleItemRow(
            business_id=business_id,
            sale_session_id=session_id,
            product_id=product_id,
            source_type="catalog",
            product_name_snapshot=name,
            quantity_input=Decimal("1"),
            unit_input="unit",
            quantity_normalized=Decimal("1"),
            unit_normalized="unit",
            unit_price=Decimal("10.00"),
            currency="MXN",
            line_total=Decimal("10.00"),
            catalog_unit_price_snapshot=Decimal("10.00"),
        )
    )
    session.flush()


def _product_sale(session, business_id, day_id) -> None:
    set_current_business_id(session, business_id)
    from app.infrastructure.persistence.models import UserRow

    actor_id = session.scalars(select(UserRow.id).where(UserRow.business_id == business_id)).one()
    product_id = _product(session, business_id, name="Secreto", price="5.00")
    _sale(
        session,
        business_id=business_id,
        actor_id=actor_id,
        day_id=day_id,
        session_id=new_uuid7(),
        confirmed_at=datetime.now(UTC),
        method="cash",
        amount="5.00",
        items=[
            _item(
                item_id=new_uuid7(),
                product_id=product_id,
                name="Secreto",
                quantity="1",
                unit="unit",
                catalog="5.00",
                price="5.00",
                total="5.00",
                created_at=datetime.now(UTC),
            )
        ],
    )


def _close_with_empty_snapshot(session, business_id, actor_id, day_id, business_date) -> None:
    set_current_business_id(session, business_id)
    day = session.get(OperationalDayRow, day_id)
    day.status = "closed"
    count_id = new_uuid7()
    closed_at = datetime.now(UTC)
    session.add(
        CashCountRow(
            id=count_id,
            business_id=business_id,
            operational_day_id=day_id,
            actor_id=actor_id,
            amount=Decimal("0.00"),
            currency="MXN",
            source="manual_capture",
            counted_at=closed_at,
        )
    )
    session.add(
        ClosingSnapshotRow(
            business_id=business_id,
            operational_day_id=day_id,
            cash_count_id=count_id,
            actor_id=actor_id,
            business_date=business_date,
            currency="MXN",
            sale_count=0,
            gross_sales_total=Decimal("0.00"),
            cash_total=Decimal("0.00"),
            card_total=Decimal("0.00"),
            transfer_total=Decimal("0.00"),
            expected_cash=Decimal("0.00"),
            counted_cash=Decimal("0.00"),
            cash_difference=Decimal("0.00"),
            cash_status="balanced",
            closed_at=closed_at,
        )
    )
    session.commit()
