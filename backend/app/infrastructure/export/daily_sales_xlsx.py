"""openpyxl adapter for one OperationalDay sales export."""

from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.application.queries.export_daily_sales import SalesExport, SalesExportLine
from app.infrastructure.export.columns import COLUMN_WIDTHS, EXPORT_COLUMNS

_MONEY_FORMAT = "0.00"
_QUANTITY_FORMAT = "0.######"
_DATE_FORMAT = "yyyy-mm-dd"
_DATETIME_FORMAT = "yyyy-mm-dd hh:mm:ss"


def render_daily_sales_xlsx(export: SalesExport) -> bytes:
    zone = ZoneInfo(export.timezone_name)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ventas"
    sheet.append(list(EXPORT_COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for line in export.lines:
        _append_line(sheet, line, zone)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:Q{sheet.max_row}"
    for index, width in enumerate(COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    fixed = datetime(1980, 1, 1, tzinfo=timezone.utc)
    workbook.properties.creator = "lumo"
    workbook.properties.lastModifiedBy = "lumo"
    workbook.properties.created = fixed
    workbook.properties.modified = fixed
    payload = BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def _append_line(sheet, line: SalesExportLine, zone: ZoneInfo) -> None:
    sheet.append([None] * len(EXPORT_COLUMNS))
    row = sheet.max_row
    local = line.sale_confirmed_at.astimezone(zone).replace(tzinfo=None, microsecond=0)
    values: list[object] = [
        line.business_date,
        str(line.sale_session_id),
        local,
        str(line.sale_item_id),
        line.product_name,
        line.source_type,
        None if line.product_id is None else str(line.product_id),
        line.quantity,
        line.unit,
        line.catalog_unit_price,
        line.unit_price,
        line.price_override_reason,
        line.line_total,
        line.currency,
        str(line.payment_id),
        line.payment_method,
        line.payment_amount,
    ]
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=column, value=value)
        header = EXPORT_COLUMNS[column - 1]
        if header in {"catalog_unit_price", "unit_price", "line_total", "payment_amount"} and value is not None:
            cell.number_format = _MONEY_FORMAT
        elif header == "quantity":
            cell.number_format = _QUANTITY_FORMAT
        elif header == "business_date":
            cell.number_format = _DATE_FORMAT
        elif header == "sale_confirmed_at":
            cell.number_format = _DATETIME_FORMAT
